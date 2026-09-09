# -*- coding: utf-8 -*-
"""
Lee y escribe directamente sobre "ECOS DATOS.xlsx" (una hoja por equipo).
Este módulo conoce la plantilla exacta de la hoja y solo toca las celdas
de datos, dejando intactos formato, colores, anchos de columna, etc.
"""
import io
import os
import re
import threading
from datetime import date, datetime
from copy import copy

import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, "ECOS DATOS.xlsx")

# Si hay credenciales de Microsoft Graph configuradas, se lee/escribe el
# Excel directamente en OneDrive en vez de en un archivo local.
USE_GRAPH = bool(os.environ.get("MS_CLIENT_ID"))
GRAPH_FILENAME = "ECOS DATOS.xlsx"
_graph_item = {"drive_id": None, "item_id": None}

_lock = threading.Lock()

COMPONENTE_LABELS = ["Power Regulator", "ACB", "PC / Computer Assembly", "Acquisition Module"]
COMPONENTE_ROWS = [10, 11, 12, 13]  # A=label(no tocar), B=version, C=maxSw, D=12nc

SONDAS_HEADER_ROW = 1
SONDAS_FIRST_ROW = 2
SONDAS_LAST_ROW = 7  # fila 8 ya es la cabecera "DICOM LOCAL"
SONDAS_COLS = {"nombre": 6, "serie": 7, "codigo12nc": 8, "pass": 9}  # F,G,H,I

DICOM_ROWS = {"ip": 9, "mask": 10, "gateway": 11, "aet": 12, "port": 13}
DICOM_VALUE_COL = 7  # G
DVDPC_ROW = 14
DVDPC_COL = 7

NOTASDICOM_ROW = 9
NOTASDICOM_COL = 9  # I

NODOS_HEADER_ROW = 19
NODOS_FIRST_ROW = 20
NODOS_LAST_ROW = 29
NODOS_COLS = {"nombre": 6, "aet": 7, "ip": 8, "puerto": 9}  # F,G,H,I

FIELD_ROWS = {
    "modelo": 1, "sn": 2, "hospital": 3, "swVersion": 4, "hwVersion": 5, "notas": 6,
}
VALUE_COL = 2  # B

FECHA_INSTALACION_ROW = 19
MANTENIMIENTO_ROWS = {
    "ultimaRevisionPM": 20, "ultimoBackup": 21, "ultimoTestGeneral": 22,
    "ultimoTestSondas": 23, "ultimoTestBaterias": 24, "borrarAvisosMtto": 25,
}

INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")

# ---------------------------------------------------------------------------
# Hojas nuevas (no forman parte de la plantilla original de cada equipo, así
# que viven en sus propias hojas para no arriesgar el formato/celdas de la
# plantilla de equipo). Cada una es una tabla simple: fila 1 = cabecera.
# ---------------------------------------------------------------------------
RESERVED_SHEETS = ("META", "HISTORIAL", "AVERIAS")

META_SHEET = "META"
META_HEADERS = ["SN", "PeriodicidadMeses", "AlertaAtrasadaEnviada"]

HISTORIAL_SHEET = "HISTORIAL"
HISTORIAL_HEADERS = ["ID", "SN", "Fecha", "Tipo", "Detalle"]

AVERIAS_SHEET = "AVERIAS"
AVERIAS_HEADERS = ["ID", "SN", "FechaApertura", "Descripcion", "Componente", "Tecnico", "Estado", "FechaResolucion", "NotasResolucion"]

MANTENIMIENTO_LABELS = {
    "ultimaRevisionPM": "Revisión PM",
    "ultimoBackup": "Backup",
    "ultimoTestGeneral": "Test general",
    "ultimoTestSondas": "Test de sondas",
    "ultimoTestBaterias": "Test de baterías",
}


def _is_real(v):
    if v is None:
        return False
    s = str(v).strip()
    return s != "" and s != "-"


def _clean(v):
    return str(v).strip() if _is_real(v) else None


def _cell_to_date_str(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return _clean(v)


def _date_str_to_value(s):
    if not s:
        return None
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        return s


def sanitize_sheet_name(sn):
    name = INVALID_SHEET_CHARS.sub("", (sn or "").strip())
    return name[:31]


def _resolve_graph_item():
    if _graph_item["drive_id"] is None:
        import graph_client
        found = graph_client.find_item_by_name(GRAPH_FILENAME)
        if not found:
            raise FileNotFoundError(f"No se encuentra '{GRAPH_FILENAME}' en el OneDrive.")
        _graph_item["drive_id"], _graph_item["item_id"], _ = found
    return _graph_item["drive_id"], _graph_item["item_id"]


def _load():
    if USE_GRAPH:
        import graph_client
        drive_id, item_id = _resolve_graph_item()
        data = graph_client.download_bytes(drive_id, item_id)
        return openpyxl.load_workbook(io.BytesIO(data))
    return openpyxl.load_workbook(EXCEL_PATH)


def _save(wb):
    if USE_GRAPH:
        import graph_client
        drive_id, item_id = _resolve_graph_item()
        buf = io.BytesIO()
        wb.save(buf)
        graph_client.upload_bytes(drive_id, item_id, buf.getvalue())
    else:
        wb.save(EXCEL_PATH)


def _parse_sheet(ws):
    sn_cell = ws.cell(row=FIELD_ROWS["sn"], column=VALUE_COL).value
    sn = _clean(sn_cell) or ws.title

    componentes = []
    for i, row in enumerate(COMPONENTE_ROWS):
        componentes.append({
            "nombre": COMPONENTE_LABELS[i],
            "version": _clean(ws.cell(row=row, column=2).value),
            "maxSw": _clean(ws.cell(row=row, column=3).value),
            "codigo12nc": _clean(ws.cell(row=row, column=4).value),
        })

    sondas = []
    for row in range(SONDAS_FIRST_ROW, SONDAS_LAST_ROW + 1):
        nombre = ws.cell(row=row, column=SONDAS_COLS["nombre"]).value
        serie = ws.cell(row=row, column=SONDAS_COLS["serie"]).value
        codigo = ws.cell(row=row, column=SONDAS_COLS["codigo12nc"]).value
        pass_ = ws.cell(row=row, column=SONDAS_COLS["pass"]).value
        if any(_is_real(v) for v in (nombre, serie, codigo, pass_)):
            sondas.append({
                "nombre": _clean(nombre), "serie": _clean(serie),
                "codigo12nc": _clean(codigo), "pass": _clean(pass_),
            })

    nodos = []
    for row in range(NODOS_FIRST_ROW, NODOS_LAST_ROW + 1):
        nombre = ws.cell(row=row, column=NODOS_COLS["nombre"]).value
        aet = ws.cell(row=row, column=NODOS_COLS["aet"]).value
        ip = ws.cell(row=row, column=NODOS_COLS["ip"]).value
        puerto = ws.cell(row=row, column=NODOS_COLS["puerto"]).value
        if any(_is_real(v) for v in (nombre, aet, ip, puerto)):
            nodos.append({
                "nombre": _clean(nombre), "aet": _clean(aet),
                "ip": _clean(ip), "puerto": _clean(puerto),
            })

    dicom = {}
    for key, row in DICOM_ROWS.items():
        dicom[key] = _clean(ws.cell(row=row, column=DICOM_VALUE_COL).value)
    dicom["dvdPc"] = _clean(ws.cell(row=DVDPC_ROW, column=DVDPC_COL).value)

    record = {
        "sn": sn,
        "modelo": _clean(ws.cell(row=FIELD_ROWS["modelo"], column=VALUE_COL).value),
        "hospital": _clean(ws.cell(row=FIELD_ROWS["hospital"], column=VALUE_COL).value),
        "swVersion": _clean(ws.cell(row=FIELD_ROWS["swVersion"], column=VALUE_COL).value),
        "hwVersion": _clean(ws.cell(row=FIELD_ROWS["hwVersion"], column=VALUE_COL).value),
        "notas": _clean(ws.cell(row=FIELD_ROWS["notas"], column=VALUE_COL).value),
        "fechaInstalacion": _cell_to_date_str(ws.cell(row=FECHA_INSTALACION_ROW, column=VALUE_COL).value),
        "notasDicom": _clean(ws.cell(row=NOTASDICOM_ROW, column=NOTASDICOM_COL).value),
        "componentes": componentes,
        "sondas": sondas,
        "nodosDicom": nodos,
        "dicom": dicom,
        "_sheet": ws.title,
    }
    for key, row in MANTENIMIENTO_ROWS.items():
        record[key] = _cell_to_date_str(ws.cell(row=row, column=VALUE_COL).value)
    return record


def list_equipos():
    with _lock:
        wb = _load()
        try:
            _, meta_map = _load_meta_map(wb)
            averias_counts = _count_averias_abiertas(wb)
            records = []
            for name in wb.sheetnames:
                if name in RESERVED_SHEETS:
                    continue
                r = _parse_sheet(wb[name])
                m = meta_map.get(r["sn"], {})
                r["periodicidadMeses"] = m.get("periodicidadMeses")
                r["averiasAbiertas"] = averias_counts.get(r["sn"], 0)
                records.append(r)
            return records
        finally:
            wb.close()


def get_equipo(sn):
    with _lock:
        wb = _load()
        try:
            if sn not in wb.sheetnames:
                return None
            r = _parse_sheet(wb[sn])
            _, meta_map = _load_meta_map(wb)
            m = meta_map.get(sn, {})
            r["periodicidadMeses"] = m.get("periodicidadMeses")
            r["averiasAbiertas"] = _count_averias_abiertas(wb).get(sn, 0)
            return r
        finally:
            wb.close()


def _clear_and_write_data(ws, record):
    ws.cell(row=FIELD_ROWS["modelo"], column=VALUE_COL).value = record.get("modelo")
    ws.cell(row=FIELD_ROWS["sn"], column=VALUE_COL).value = record.get("sn")
    ws.cell(row=FIELD_ROWS["hospital"], column=VALUE_COL).value = record.get("hospital")
    ws.cell(row=FIELD_ROWS["swVersion"], column=VALUE_COL).value = record.get("swVersion")
    ws.cell(row=FIELD_ROWS["hwVersion"], column=VALUE_COL).value = record.get("hwVersion")
    ws.cell(row=FIELD_ROWS["notas"], column=VALUE_COL).value = record.get("notas")

    ws.cell(row=FECHA_INSTALACION_ROW, column=VALUE_COL).value = _date_str_to_value(record.get("fechaInstalacion"))
    for key, row in MANTENIMIENTO_ROWS.items():
        ws.cell(row=row, column=VALUE_COL).value = _date_str_to_value(record.get(key))

    componentes = record.get("componentes") or []
    for i, row in enumerate(COMPONENTE_ROWS):
        c = componentes[i] if i < len(componentes) else {}
        ws.cell(row=row, column=2).value = c.get("version")
        ws.cell(row=row, column=3).value = c.get("maxSw")
        ws.cell(row=row, column=4).value = c.get("codigo12nc")

    for row in range(SONDAS_FIRST_ROW, SONDAS_LAST_ROW + 1):
        for col in SONDAS_COLS.values():
            ws.cell(row=row, column=col).value = None
    sondas = (record.get("sondas") or [])[:SONDAS_LAST_ROW - SONDAS_FIRST_ROW + 1]
    for i, s in enumerate(sondas):
        row = SONDAS_FIRST_ROW + i
        ws.cell(row=row, column=SONDAS_COLS["nombre"]).value = s.get("nombre")
        ws.cell(row=row, column=SONDAS_COLS["serie"]).value = s.get("serie")
        ws.cell(row=row, column=SONDAS_COLS["codigo12nc"]).value = s.get("codigo12nc")
        ws.cell(row=row, column=SONDAS_COLS["pass"]).value = s.get("pass")

    dicom = record.get("dicom") or {}
    for key, row in DICOM_ROWS.items():
        cell = ws.cell(row=row, column=DICOM_VALUE_COL)
        cell.value = dicom.get(key)
        cell.number_format = "@"
    ws.cell(row=DVDPC_ROW, column=DVDPC_COL).value = dicom.get("dvdPc")

    ws.cell(row=NOTASDICOM_ROW, column=NOTASDICOM_COL).value = record.get("notasDicom")

    for row in range(NODOS_FIRST_ROW, NODOS_LAST_ROW + 1):
        for col in NODOS_COLS.values():
            ws.cell(row=row, column=col).value = None
    nodos = (record.get("nodosDicom") or [])[:NODOS_LAST_ROW - NODOS_FIRST_ROW + 1]
    for i, n in enumerate(nodos):
        row = NODOS_FIRST_ROW + i
        ws.cell(row=row, column=NODOS_COLS["nombre"]).value = n.get("nombre")
        ws.cell(row=row, column=NODOS_COLS["aet"]).value = n.get("aet")
        ws.cell(row=row, column=NODOS_COLS["ip"]).value = n.get("ip")
        ws.cell(row=row, column=NODOS_COLS["puerto"]).value = n.get("puerto")


def _new_sheet_from_template(wb, sheet_name):
    template = None
    for name in wb.sheetnames:
        if name not in RESERVED_SHEETS:
            template = wb[name]
            break
    if template is None:
        raise RuntimeError("No hay ninguna plantilla de equipo en el Excel para copiar.")
    ws = wb.copy_worksheet(template)
    ws.title = sheet_name
    for i, row in enumerate(COMPONENTE_ROWS):
        ws.cell(row=row, column=1).value = COMPONENTE_LABELS[i]
    return ws


class SheetExistsError(Exception):
    pass


class SheetNotFoundError(Exception):
    pass


class AveriaNotFoundError(Exception):
    pass


def create_equipo(record):
    sn = sanitize_sheet_name(record.get("sn"))
    if not sn:
        raise ValueError("Número de serie vacío.")
    with _lock:
        wb = _load()
        try:
            if sn in wb.sheetnames:
                raise SheetExistsError(sn)
            ws = _new_sheet_from_template(wb, sn)
            record = dict(record, sn=sn)
            _clear_and_write_data(ws, record)
            _set_periodicidad(wb, sn, record.get("periodicidadMeses"))
            _log_historial(wb, sn, "creacion", "Equipo dado de alta.")
            _save(wb)
        finally:
            wb.close()
    return get_equipo(sn)


def update_equipo(original_sn, record):
    new_sn = sanitize_sheet_name(record.get("sn")) or original_sn
    with _lock:
        wb = _load()
        try:
            if original_sn not in wb.sheetnames:
                raise SheetNotFoundError(original_sn)
            if new_sn != original_sn and new_sn in wb.sheetnames:
                raise SheetExistsError(new_sn)
            ws = wb[original_sn]
            before = _parse_sheet(ws)
            record = dict(record, sn=new_sn)
            _clear_and_write_data(ws, record)
            if new_sn != original_sn:
                ws.title = new_sn
                _rename_sn_everywhere(wb, original_sn, new_sn)
            _set_periodicidad(wb, new_sn, record.get("periodicidadMeses"))
            _log_mantenimiento_changes(wb, new_sn, before, record)
            _save(wb)
        finally:
            wb.close()
    return get_equipo(new_sn)


def delete_equipo(sn):
    with _lock:
        wb = _load()
        try:
            if sn not in wb.sheetnames:
                raise SheetNotFoundError(sn)
            _log_historial(wb, sn, "eliminacion", "Equipo eliminado del inventario.")
            del wb[sn]
            _save(wb)
        finally:
            wb.close()


# ---------------------------------------------------------------------------
# META (periodicidad de revisión por equipo + estado de la última alerta)
# ---------------------------------------------------------------------------

def _sheet_rows(ws):
    """Itera (fila, [valores]) para cada fila de datos no vacía (salta la cabecera)."""
    for row in range(2, ws.max_row + 1):
        values = [ws.cell(row=row, column=c).value for c in range(1, ws.max_column + 1)]
        if any(_is_real(v) for v in values):
            yield row, values


def _get_or_create_sheet(wb, name, headers):
    if name in wb.sheetnames:
        return wb[name]
    ws = wb.create_sheet(name)
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i).value = h
    return ws


def _next_id(ws):
    max_id = 0
    for _, values in _sheet_rows(ws):
        try:
            max_id = max(max_id, int(values[0]))
        except (TypeError, ValueError):
            pass
    return max_id + 1


def _load_meta_map(wb):
    ws = _get_or_create_sheet(wb, META_SHEET, META_HEADERS)
    out = {}
    for row, values in _sheet_rows(ws):
        sn = _clean(values[0])
        if not sn:
            continue
        out[sn] = {
            "row": row,
            "periodicidadMeses": _clean(values[1]) if len(values) > 1 else None,
            "alertaAtrasadaEnviada": _clean(values[2]) if len(values) > 2 else None,
        }
    return ws, out


def _get_meta_row(ws, meta_map, sn):
    if sn in meta_map:
        return meta_map[sn]["row"]
    row = ws.max_row + 1
    ws.cell(row=row, column=1).value = sn
    meta_map[sn] = {"row": row, "periodicidadMeses": None, "alertaAtrasadaEnviada": None}
    return row


def _set_periodicidad(wb, sn, value):
    ws, meta_map = _load_meta_map(wb)
    row = _get_meta_row(ws, meta_map, sn)
    ws.cell(row=row, column=2).value = value


def _rename_sn_everywhere(wb, old_sn, new_sn):
    ws_meta, meta_map = _load_meta_map(wb)
    if old_sn in meta_map:
        ws_meta.cell(row=meta_map[old_sn]["row"], column=1).value = new_sn
    for sheet_name in (HISTORIAL_SHEET, AVERIAS_SHEET):
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row, values in _sheet_rows(ws):
                if _clean(values[1]) == old_sn:
                    ws.cell(row=row, column=2).value = new_sn


def get_alertas_estado():
    """dict SN -> True si ya se avisó de que está atrasado (para no repetir el aviso)."""
    with _lock:
        wb = _load()
        try:
            _, meta_map = _load_meta_map(wb)
            return {sn: (str(v.get("alertaAtrasadaEnviada") or "").upper() == "SI") for sn, v in meta_map.items()}
        finally:
            wb.close()


def set_alertas_estado(estado_por_sn):
    """estado_por_sn: dict SN -> True/False. Persiste qué equipos ya han generado aviso."""
    with _lock:
        wb = _load()
        try:
            ws, meta_map = _load_meta_map(wb)
            for sn, alertado in estado_por_sn.items():
                row = _get_meta_row(ws, meta_map, sn)
                ws.cell(row=row, column=3).value = "SI" if alertado else "NO"
            _save(wb)
        finally:
            wb.close()


# ---------------------------------------------------------------------------
# HISTORIAL (registro de eventos por equipo: creación, revisiones, notas...)
# ---------------------------------------------------------------------------

def _log_historial(wb, sn, tipo, detalle, fecha=None):
    ws = _get_or_create_sheet(wb, HISTORIAL_SHEET, HISTORIAL_HEADERS)
    new_id = _next_id(ws)
    row = ws.max_row + 1
    ws.cell(row=row, column=1).value = new_id
    ws.cell(row=row, column=2).value = sn
    ws.cell(row=row, column=3).value = fecha or date.today()
    ws.cell(row=row, column=4).value = tipo
    ws.cell(row=row, column=5).value = detalle
    return new_id


def _log_mantenimiento_changes(wb, sn, before, after):
    for key, label in MANTENIMIENTO_LABELS.items():
        old_v = before.get(key)
        new_v = after.get(key)
        if new_v and new_v != old_v:
            _log_historial(wb, sn, "mantenimiento", label + " registrada: " + new_v)


def list_historial(sn):
    with _lock:
        wb = _load()
        try:
            if HISTORIAL_SHEET not in wb.sheetnames:
                return []
            ws = wb[HISTORIAL_SHEET]
            items = []
            for _, values in _sheet_rows(ws):
                if _clean(values[1]) != sn:
                    continue
                items.append({
                    "id": values[0], "sn": values[1],
                    "fecha": _cell_to_date_str(values[2]),
                    "tipo": _clean(values[3]), "detalle": _clean(values[4]),
                })
            items.sort(key=lambda x: x["fecha"] or "", reverse=True)
            return items
        finally:
            wb.close()


def add_historial_nota(sn, detalle):
    if not (detalle or "").strip():
        raise ValueError("La nota está vacía.")
    with _lock:
        wb = _load()
        try:
            if sn not in wb.sheetnames:
                raise SheetNotFoundError(sn)
            _log_historial(wb, sn, "nota", detalle.strip())
            _save(wb)
        finally:
            wb.close()
    return list_historial(sn)


# ---------------------------------------------------------------------------
# AVERIAS (incidencias por equipo, abiertas/cerradas)
# ---------------------------------------------------------------------------

def list_averias(sn):
    with _lock:
        wb = _load()
        try:
            if AVERIAS_SHEET not in wb.sheetnames:
                return []
            ws = wb[AVERIAS_SHEET]
            items = []
            for _, values in _sheet_rows(ws):
                if _clean(values[1]) != sn:
                    continue
                items.append({
                    "id": values[0], "sn": values[1],
                    "fechaApertura": _cell_to_date_str(values[2]),
                    "descripcion": _clean(values[3]),
                    "componente": _clean(values[4]) if len(values) > 4 else None,
                    "tecnico": _clean(values[5]) if len(values) > 5 else None,
                    "estado": (_clean(values[6]) if len(values) > 6 else None) or "Abierta",
                    "fechaResolucion": _cell_to_date_str(values[7]) if len(values) > 7 else None,
                    "notasResolucion": _clean(values[8]) if len(values) > 8 else None,
                })
            items.sort(key=lambda x: x["fechaApertura"] or "", reverse=True)
            return items
        finally:
            wb.close()


def _count_averias_abiertas(wb):
    if AVERIAS_SHEET not in wb.sheetnames:
        return {}
    ws = wb[AVERIAS_SHEET]
    counts = {}
    for _, values in _sheet_rows(ws):
        sn = _clean(values[1]) if len(values) > 1 else None
        estado = (_clean(values[6]) if len(values) > 6 else None) or "Abierta"
        if sn and estado == "Abierta":
            counts[sn] = counts.get(sn, 0) + 1
    return counts


def create_averia(sn, record):
    descripcion = (record.get("descripcion") or "").strip()
    if not descripcion:
        raise ValueError("Falta la descripción de la avería.")
    with _lock:
        wb = _load()
        try:
            if sn not in wb.sheetnames:
                raise SheetNotFoundError(sn)
            ws = _get_or_create_sheet(wb, AVERIAS_SHEET, AVERIAS_HEADERS)
            new_id = _next_id(ws)
            row = ws.max_row + 1
            ws.cell(row=row, column=1).value = new_id
            ws.cell(row=row, column=2).value = sn
            ws.cell(row=row, column=3).value = _date_str_to_value(record.get("fechaApertura")) or date.today()
            ws.cell(row=row, column=4).value = descripcion
            ws.cell(row=row, column=5).value = record.get("componente")
            ws.cell(row=row, column=6).value = record.get("tecnico")
            ws.cell(row=row, column=7).value = "Abierta"
            _log_historial(wb, sn, "averia_abierta", "Avería registrada: " + descripcion)
            _save(wb)
        finally:
            wb.close()
    return list_averias(sn)


def update_averia(sn, averia_id, record):
    with _lock:
        wb = _load()
        try:
            if AVERIAS_SHEET not in wb.sheetnames:
                raise AveriaNotFoundError(str(averia_id))
            ws = wb[AVERIAS_SHEET]
            target_row = None
            for row, values in _sheet_rows(ws):
                if str(values[0]) == str(averia_id) and _clean(values[1]) == sn:
                    target_row = row
                    break
            if target_row is None:
                raise AveriaNotFoundError(str(averia_id))
            if "descripcion" in record and record.get("descripcion"):
                ws.cell(row=target_row, column=4).value = record.get("descripcion")
            if "componente" in record:
                ws.cell(row=target_row, column=5).value = record.get("componente")
            if "tecnico" in record:
                ws.cell(row=target_row, column=6).value = record.get("tecnico")
            estado = record.get("estado")
            if estado:
                ws.cell(row=target_row, column=7).value = estado
            if estado == "Cerrada":
                ws.cell(row=target_row, column=8).value = _date_str_to_value(record.get("fechaResolucion")) or date.today()
                ws.cell(row=target_row, column=9).value = record.get("notasResolucion")
                detalle = "Avería resuelta"
                if record.get("notasResolucion"):
                    detalle += ": " + record.get("notasResolucion")
                _log_historial(wb, sn, "averia_cerrada", detalle)
            _save(wb)
        finally:
            wb.close()
    return list_averias(sn)
